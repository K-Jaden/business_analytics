#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
SHAP 채널 선택 — '평균 중요도 임계값(기준 점선)' (1장)
발표자료 v3 SLIDE21 "SHAP이 밝힌 것" 의 구조·색·★·느낌을 그대로 가져오고,
'평균 임계값(기준 점선)' 요소를 추가한 버전.
값 = Model A SHAP (results/shap_summary.csv, 21쪽과 동일).
선택 = 채널 평균 이상(above-mean) → ★ 표시. DM = results/dm_fixed_selection.csv.
Usage : python scripts/make_slide_shap_threshold.py
Output: slide_shap_threshold_v2.pptx
"""
import os
import io
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(BASE, "results")

for f in ["Malgun Gothic", "맑은 고딕"]:
    if any(f.lower() in fn.name.lower() for fn in fm.fontManager.ttflist):
        plt.rcParams["font.family"] = f
        break
plt.rcParams["axes.unicode_minus"] = False

NAVY = RGBColor(0x1F, 0x2D, 0x4E)
BLUE = RGBColor(0x2E, 0x6D, 0xA4)
ORANGE = RGBColor(0xE8, 0x77, 0x22)
GREEN = RGBColor(0x1E, 0x7E, 0x34)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DGRAY = RGBColor(0x44, 0x44, 0x44)
LGRAY = RGBColor(0xF4, 0xF6, 0xF9)
INK = RGBColor(0x0A, 0x0A, 0x0A)

CH = ["CH1", "CH2", "CH3", "CH4", "CH5"]
CH_KEY = {"CH1": "score_ch1", "CH2": "score_ch2", "CH3": "score_ch3",
          "CH4": "score_ch4", "CH5": "score_ch5"}
CH_COLOR = {"CH1": "#1F8FE0", "CH2": "#E08A20", "CH3": "#9AA0A6",
            "CH4": "#E8B62C", "CH5": "#C0392B"}
CH_NAME = {"CH1": "GT(검색량)", "CH2": "뉴스량", "CH3": "뉴스감성",
           "CH4": "YT조회", "CH5": "YT감성"}
ASSET_KEY = {"스니커즈": "sneakers", "카드": "cards", "레고": "lego"}
ASSETS = ["스니커즈", "카드", "레고"]

# ── Model A SHAP 값 (results/shap_summary.csv = 21쪽과 동일) ────────────────
_sh = pd.read_csv(os.path.join(RES, "shap_summary.csv"))
_sh = _sh[_sh["model"] == "Model A"]
SHAP = {}
for kr, a in ASSET_KEY.items():
    sub = _sh[_sh["asset_type"] == a].set_index("feature")["mean_abs_shap"]
    SHAP[kr] = {ch: float(sub.get(CH_KEY[ch], 0.0)) for ch in CH}

# 평균 임계값(above-mean) 선택
SELECTED = {}
for kr in ASSETS:
    d = SHAP[kr]
    m = sum(d.values()) / len(CH)
    SELECTED[kr] = [ch for ch in sorted(CH, key=lambda c: d[c], reverse=True) if d[ch] >= m]

# DM 결과 (고정 선택 기준)
_dm = pd.read_csv(os.path.join(RES, "dm_fixed_selection.csv")).set_index("asset_type")
DM_P = {kr: float(_dm.loc[a, "dm_p"]) for kr, a in ASSET_KEY.items()}

# ════════════════════════════════════════════════════════════════════════
# 그래프 — 21쪽 스타일(채널별 색 + ★ + 색-이름 범례) + 기준 점선(평균)
# ════════════════════════════════════════════════════════════════════════
fig, axes = plt.subplots(1, 3, figsize=(9.6, 3.05), dpi=200)
for ax, asset in zip(axes, ASSETS):
    d = SHAP[asset]
    vals = [d[c] for c in CH]
    mx = max(vals) if max(vals) > 0 else 1.0
    mean_imp = sum(vals) / len(CH)
    sel_set = set(SELECTED[asset])
    xpos = np.arange(len(CH))
    for xi, c in zip(xpos, CH):
        sel = c in sel_set
        ax.bar(xi, d[c], width=0.72, color=CH_COLOR[c],
               alpha=1.0 if sel else 0.32, edgecolor="white", linewidth=0.6, zorder=3)
        ax.text(xi, d[c] + mx * 0.03, f"{d[c]:.3f}", ha="center", va="bottom",
                fontsize=7, fontweight="bold" if sel else "normal",
                color="#222222" if sel else "#9AA0A6")
        if sel:
            ax.text(xi, d[c] + mx * 0.14, "★", ha="center", va="bottom",
                    fontsize=10.5, color="#2E6DA4", zorder=4)
    ax.axhline(mean_imp, color="#E8730F", linestyle="--", linewidth=1.8, zorder=5)
    ax.text(2, mean_imp + mx * 0.04, f"기준선 {mean_imp:.3f}",
            color="#E8730F", fontsize=7, fontweight="bold", ha="center", va="bottom", zorder=6)
    ax.set_xticks(xpos)
    ax.set_xticklabels(CH, fontsize=8.5)
    ax.set_title(f"{asset}  →  {'+'.join(SELECTED[asset])}",
                 fontsize=10.5, fontweight="bold", pad=14, color="#1F2D4E")
    ax.set_ylim(0, mx * 1.36)
    ax.set_xlim(-0.7, len(CH) - 0.3)
    ax.spines[["top", "right"]].set_visible(False)
    ax.tick_params(axis="y", labelsize=7)
# 색-이름 범례 (21쪽 느낌)
from matplotlib.patches import Patch
handles = [Patch(facecolor=CH_COLOR[c], label=f"{c} {CH_NAME[c]}") for c in CH]
fig.legend(handles=handles, loc="lower center", ncol=5, fontsize=8.5,
           frameon=False, bbox_to_anchor=(0.5, -0.04))
fig.suptitle("SHAP 채널 중요도 (Model A 기준)  ·  ★ = 기준 점선(채널 평균) 이상 → 선택",
             fontsize=11, fontweight="bold", y=1.04, color="#1F2D4E")
plt.tight_layout(rect=[0, 0.06, 1, 1])
fig_buf = io.BytesIO()
fig.savefig(fig_buf, dpi=195, bbox_inches="tight", facecolor="white")
plt.close(fig)

# ════════════════════════════════════════════════════════════════════════
# 슬라이드
# ════════════════════════════════════════════════════════════════════════
prs = Presentation()
prs.slide_width = Inches(13.33)
prs.slide_height = Inches(7.5)


def box(sld, l, t, w, h, fill=None, line=None, line_w=1.0, round_=False):
    shp = sld.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(l), Inches(t), Inches(w), Inches(h))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line; shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    return shp


def text(sld, l, t, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, space=2):
    tb = sld.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = Pt(4); tf.margin_right = Pt(4)
    tf.margin_top = Pt(1); tf.margin_bottom = Pt(1)
    for i, para in enumerate(runs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align; p.space_after = Pt(space)
        for (txt, size, bold, color) in para:
            r = p.add_run(); r.text = txt
            r.font.size = Pt(size); r.font.bold = bold
            r.font.color.rgb = color; r.font.name = "맑은 고딕"
    return tb


def header(sld, title, subtitle_runs):
    box(sld, 0, 0, 13.33, 0.82, fill=NAVY)
    text(sld, 0.25, 0.06, 12.8, 0.70, [[(title, 21, True, WHITE)]], anchor=MSO_ANCHOR.MIDDLE)
    box(sld, 0, 0.82, 13.33, 0.045, fill=BLUE)
    text(sld, 0.25, 0.92, 12.8, 0.40, [subtitle_runs], anchor=MSO_ANCHOR.MIDDLE)


s1 = prs.slides.add_slide(prs.slide_layouts[6])
header(s1, "SHAP이 밝힌 것 — 채널 선택 기준은 '평균(기준 점선)'",
       [("채널 5개 SHAP의 평균에 점선을 긋고 ", 13, False, DGRAY),
        ("그 위(★)만 선택", 13, True, ORANGE),
        (" — 반 평균 점수와 같은 원리", 13, False, DGRAY)])

MARGIN = 0.30

# ── 그래프 (상단, 전폭) ────────────────────────────────────────────────────
FIG_T, FIG_W, FIG_H = 1.32, 13.33 - 2 * MARGIN, 3.66
box(s1, MARGIN, FIG_T, FIG_W, FIG_H, fill=LGRAY, line=RGBColor(0xCF, 0xD6, 0xDF), line_w=1.0, round_=True)
from PIL import Image as _PILImage
fig_buf.seek(0)
_w_px, _h_px = _PILImage.open(fig_buf).size
ratio = _h_px / _w_px
pic_w = FIG_W - 0.30
pic_h = pic_w * ratio
if pic_h > FIG_H - 0.20:
    pic_h = FIG_H - 0.20
    pic_w = pic_h / ratio
fig_buf.seek(0)
s1.shapes.add_picture(fig_buf, Inches(MARGIN + (FIG_W - pic_w) / 2), Inches(FIG_T + 0.10),
                      width=Inches(pic_w))

# ── 하단: 자산별 박스 3개 (21쪽 구조) ──────────────────────────────────────
BOX_T = FIG_T + FIG_H + 0.16
BOX_H = 1.62
GAP = 0.20
BOX_W = (13.33 - 2 * MARGIN - 2 * GAP) / 3
ASSET_BAR = {"스니커즈": BLUE, "카드": GREEN, "레고": RGBColor(0x7B, 0x3F, 0xA0)}
CH_FULL = {"CH1": "Google Trends", "CH2": "뉴스 보도량", "CH3": "뉴스 감성",
           "CH4": "YT 조회수", "CH5": "YT 댓글감성"}
for i, asset in enumerate(ASSETS):
    x = MARGIN + i * (BOX_W + GAP)
    box(s1, x, BOX_T, BOX_W, BOX_H, fill=LGRAY, line=RGBColor(0xCF, 0xD6, 0xDF), line_w=1.0, round_=True)
    box(s1, x, BOX_T, BOX_W, 0.40, fill=ASSET_BAR[asset], round_=True)
    text(s1, x, BOX_T + 0.02, BOX_W, 0.36,
         [[(f"{asset}  →  선택 {len(SELECTED[asset])}개", 12, True, WHITE)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    runs = []
    for rank, ch in enumerate(SELECTED[asset], 1):
        runs.append([("★ ", 10.5, True, BLUE), (f"{ch} {CH_FULL[ch]}", 10.5, True, NAVY),
                     (f"  ({SHAP[asset][ch]:.3f})", 10, False, DGRAY)])
    p = DM_P[asset]
    mark = " (경계)" if 0.05 < p < 0.10 else ""
    runs.append([("DM vs 전채널  p=", 9.5, False, INK), (f"{p:.3f}", 9.5, True, GREEN),
                 (f" → 동등{mark}", 9.5, False, INK)])
    text(s1, x + 0.16, BOX_T + 0.48, BOX_W - 0.32, BOX_H - 0.56, runs, space=5)

# ── 출처 + 결론 ────────────────────────────────────────────────────────────
SRC_T = BOX_T + BOX_H + 0.10
text(s1, MARGIN + 0.04, SRC_T, 13.33 - 2 * MARGIN, 0.26,
     [[("※ 값 = Model A SHAP 기여도(전채널 동시 투입) · 선택 기준 = ", 9.5, False, DGRAY),
       ("평균 중요도 임계값 (scikit-learn SelectFromModel 기본값 threshold='mean')", 9.5, True, BLUE),
       (" · Wang et al. (2024), J. Big Data 11:44", 9.5, False, DGRAY)]])

BAN_T = 7.5 - 0.52
box(s1, MARGIN, BAN_T, 13.33 - 2 * MARGIN, 0.40, fill=NAVY)
text(s1, MARGIN + 0.15, BAN_T, 13.33 - 2 * MARGIN - 0.3, 0.40,
     [[("결론  ", 12.5, True, WHITE),
       ("'평균(기준 점선)보다 기여가 큰 채널만 남긴다'", 11.5, True, WHITE),
       ("는 객관적 규칙 → 선별 모델이 전채널과 통계적으로 동등 → 간결성 입증", 11.5, False, WHITE)]],
     anchor=MSO_ANCHOR.MIDDLE)

for cand in ["slide_shap_threshold_v2.pptx", "slide_shap_threshold_v3.pptx",
             "slide_shap_threshold_v4.pptx"]:
    try:
        prs.save(os.path.join(BASE, cand))
        print(f"[OK] {cand}  (slides={len(prs.slides._sldIdLst)})")
        break
    except PermissionError:
        print(f"  (잠김: {cand} → 다음 이름)")
print("SHAP:", {k: {c: round(v, 3) for c, v in d.items()} for k, d in SHAP.items()})
print("SELECTED:", SELECTED)
print("DM_P:", DM_P)
